"""
============================================================
DocuForge AI
Production Offline Dataset Generator
============================================================

Purpose:
    Extract labelled training samples from properly formatted
    DOCX documents.

Input:
    backend/ml/training_documents/*.docx

Output:
    backend/dataset/document_elements.csv

Offline:
    YES
    No internet
    No cloud API
    No external AI

Labels:
    Title
    Chapter Title
    Heading 1
    Heading 2
    Heading 3
    Paragraph
    List
    Table
    Figure
    Table Caption
    Figure Caption
    Reference

The generated dataset contains:

    Text features
    Word style
    Font size
    Bold
    Italic
    Underline
    Alignment
    Indentation
    Paragraph spacing
    Line spacing
    Position
    Text statistics
    Text pattern features
    Image information
    Table information
    Source document

============================================================
"""

from pathlib import Path
from typing import Optional
import csv
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# PATHS
# ============================================================

# Current file:
# backend/ml/dataset_generator.py

ML_DIR = Path(__file__).resolve().parent

# backend/
BACKEND_DIR = ML_DIR.parent

# Input:
# backend/ml/training_documents/
INPUT_DIR = ML_DIR / "training_documents"

# Dataset:
# backend/dataset/
DATASET_DIR = BACKEND_DIR / "dataset"

# Final CSV:
# backend/dataset/document_elements.csv
OUTPUT_FILE = DATASET_DIR / "document_elements.csv"


# ============================================================
# VALID LABELS
# ============================================================

VALID_LABELS = [
    "Title",
    "Chapter Title",
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Paragraph",
    "List",
    "Table",
    "Figure",
    "Table Caption",
    "Figure Caption",
    "Reference",
]


# ============================================================
# CSV COLUMNS
# ============================================================

CSV_FIELDS = [

    # Basic information
    "text",
    "label",

    # Source
    "source_document",

    # Word style
    "style",
    "style_id",

    # Font
    "font_size",
    "bold",
    "italic",
    "underline",

    # Paragraph formatting
    "alignment",
    "left_indent",
    "right_indent",
    "first_line_indent",
    "space_before",
    "space_after",
    "line_spacing",

    # Position
    "paragraph_index",

    # Text statistics
    "text_length",
    "word_count",
    "uppercase_ratio",
    "digit_ratio",

    # Text pattern features
    "starts_with_number",
    "starts_with_chapter",
    "starts_with_figure",
    "starts_with_table",
    "looks_like_reference",

    # Document element information
    "has_image",
    "has_table",
]


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(text: str) -> str:
    """
    Clean paragraph text.
    """

    if not text:
        return ""

    text = text.replace("\n", " ")
    text = text.replace("\t", " ")

    # Remove repeated spaces
    text = re.sub(r"\s+", " ", text)

    return text.strip()


# ============================================================
# SAFE NUMBER CONVERSION
# ============================================================

def safe_float(value) -> float:

    try:
        return float(value)

    except (TypeError, ValueError):

        return 0.0


# ============================================================
# FONT SIZE
# ============================================================

def get_font_size(paragraph) -> float:
    """
    Get the first explicitly defined font size.

    If different runs have different sizes,
    the first explicit size is used.
    """

    for run in paragraph.runs:

        if run.font.size is not None:

            try:

                return round(
                    run.font.size.pt,
                    2
                )

            except Exception:

                pass

    return 0.0


# ============================================================
# BOLD
# ============================================================

def get_bold(paragraph) -> int:

    for run in paragraph.runs:

        if run.bold:

            return 1

    return 0


# ============================================================
# ITALIC
# ============================================================

def get_italic(paragraph) -> int:

    for run in paragraph.runs:

        if run.italic:

            return 1

    return 0


# ============================================================
# UNDERLINE
# ============================================================

def get_underline(paragraph) -> int:

    for run in paragraph.runs:

        if run.underline:

            return 1

    return 0


# ============================================================
# ALIGNMENT
# ============================================================

def get_alignment(paragraph) -> str:
    """
    Convert Word alignment into readable text.
    """

    alignment = paragraph.alignment

    if alignment == WD_ALIGN_PARAGRAPH.LEFT:

        return "left"

    if alignment == WD_ALIGN_PARAGRAPH.CENTER:

        return "center"

    if alignment == WD_ALIGN_PARAGRAPH.RIGHT:

        return "right"

    if alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:

        return "justify"

    return "unknown"


# ============================================================
# INDENTATION
# ============================================================

def get_indent(
    paragraph,
    attribute: str
) -> float:

    try:

        value = getattr(
            paragraph.paragraph_format,
            attribute,
            None
        )

        if value is None:

            return 0.0

        return round(
            value.pt,
            2
        )

    except Exception:

        return 0.0


# ============================================================
# SPACING
# ============================================================

def get_spacing(
    paragraph,
    attribute: str
) -> float:

    try:

        value = getattr(
            paragraph.paragraph_format,
            attribute,
            None
        )

        if value is None:

            return 0.0

        return round(
            value.pt,
            2
        )

    except Exception:

        return 0.0


# ============================================================
# LINE SPACING
# ============================================================

def get_line_spacing(paragraph) -> float:
    """
    Extract Word line spacing.

    Examples:
        1.0
        1.15
        1.5
        2.0
    """

    try:

        value = (
            paragraph
            .paragraph_format
            .line_spacing
        )

        if value is None:

            return 0.0

        # Multiple
        if isinstance(value, (int, float)):

            return round(
                float(value),
                2
            )

        # Length
        if hasattr(value, "pt"):

            return round(
                value.pt,
                2
            )

        return safe_float(value)

    except Exception:

        return 0.0


# ============================================================
# IMAGE DETECTION
# ============================================================

def paragraph_contains_image(
    paragraph
) -> bool:
    """
    Detect images inside a paragraph.
    """

    try:

        xml = paragraph._p.xml

        return (
            "<w:drawing" in xml
            or "<pic:pic" in xml
            or "<v:shape" in xml
        )

    except Exception:

        return False


# ============================================================
# TEXT PATTERNS
# ============================================================

def starts_with_number(
    text: str
) -> int:
    """
    Detect:

        1 Introduction
        1. Introduction
        1.1 Background
        2.3.1 Methodology
        1) Introduction
        1 - Introduction
    """

    pattern = (
        r"^\s*\d+"
        r"(?:\.\d+)*"
        r"(?:[\s\.\)\-:])"
    )

    return int(
        bool(
            re.match(
                pattern,
                text
            )
        )
    )


# ============================================================
# CHAPTER PATTERN
# ============================================================

def starts_with_chapter(
    text: str
) -> int:

    pattern = r"^\s*chapter\s+\d+"

    return int(
        bool(
            re.match(
                pattern,
                text,
                re.IGNORECASE
            )
        )
    )


# ============================================================
# FIGURE PATTERN
# ============================================================

def starts_with_figure(
    text: str
) -> int:

    pattern = (
        r"^\s*"
        r"(?:figure|fig\.?)"
        r"\s*\d+"
    )

    return int(
        bool(
            re.match(
                pattern,
                text,
                re.IGNORECASE
            )
        )
    )


# ============================================================
# TABLE PATTERN
# ============================================================

def starts_with_table(
    text: str
) -> int:

    pattern = (
        r"^\s*"
        r"table"
        r"\s*\d+"
    )

    return int(
        bool(
            re.match(
                pattern,
                text,
                re.IGNORECASE
            )
        )
    )


# ============================================================
# REFERENCE DETECTION
# ============================================================

def looks_like_reference(
    text: str
) -> int:
    """
    Detect common reference patterns.

    Examples:

        [1] Smith, J.

        [12] Kumar et al.

        1. Smith et al.

        Smith, J. (2024)

        doi:10.xxxx

        https://...

        et al.
    """

    patterns = [

        # [1]
        r"^\s*\[\d+\]",

        # 1. Smith
        r"^\s*\d+\.\s+[A-Z][a-z]+",

        # DOI
        r"\bdoi\s*:",

        # URL
        r"\bhttps?://",

        # (2024)
        r"\(\s*\d{4}\s*\)",

        # et al.
        r"\bet\s+al\.",

    ]

    for pattern in patterns:

        if re.search(
            pattern,
            text,
            re.IGNORECASE
        ):

            return 1

    return 0


# ============================================================
# UPPERCASE RATIO
# ============================================================

def calculate_uppercase_ratio(
    text: str
) -> float:

    letters = [
        char
        for char in text
        if char.isalpha()
    ]

    if not letters:

        return 0.0

    uppercase = [
        char
        for char in letters
        if char.isupper()
    ]

    return round(
        len(uppercase) /
        len(letters),
        3
    )


# ============================================================
# DIGIT RATIO
# ============================================================

def calculate_digit_ratio(
    text: str
) -> float:

    if not text:

        return 0.0

    digits = sum(
        char.isdigit()
        for char in text
    )

    return round(
        digits / len(text),
        3
    )


# ============================================================
# WORD STYLE CLASSIFICATION
# ============================================================

def classify_by_word_style(
    paragraph
) -> Optional[str]:
    """
    First-level classification.

    Word styles are strong signals because properly
    formatted training documents contain meaningful
    paragraph styles.
    """

    try:

        style_name = (
            paragraph.style.name
            or ""
        ).strip()

    except Exception:

        style_name = ""

    style_lower = (
        style_name.lower()
    )

    text = clean_text(
        paragraph.text
    )

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    if style_lower == "title":

        return "Title"

    # --------------------------------------------------------
    # CHAPTER TITLE
    # --------------------------------------------------------

    if (
        "chapter" in style_lower
        and "title" in style_lower
    ):

        return "Chapter Title"

    # --------------------------------------------------------
    # HEADING 1
    # --------------------------------------------------------

    if style_lower in (
        "heading 1",
        "heading1",
    ):

        # If explicitly Chapter X
        if starts_with_chapter(text):

            return "Chapter Title"

        # Short Heading 1 may represent
        # chapter title in books/theses.
        if len(text.split()) <= 8:

            return "Chapter Title"

        return "Heading 1"

    # --------------------------------------------------------
    # HEADING 2
    # --------------------------------------------------------

    if style_lower in (
        "heading 2",
        "heading2",
    ):

        return "Heading 2"

    # --------------------------------------------------------
    # HEADING 3
    # --------------------------------------------------------

    if style_lower in (
        "heading 3",
        "heading3",
    ):

        return "Heading 3"

    # --------------------------------------------------------
    # CAPTION
    # --------------------------------------------------------

    if "caption" in style_lower:

        if starts_with_figure(text):

            return "Figure Caption"

        if starts_with_table(text):

            return "Table Caption"

        # Generic caption
        return "Figure Caption"

    # --------------------------------------------------------
    # LIST
    # --------------------------------------------------------

    if (
        "list" in style_lower
        or "bullet" in style_lower
        or "number" in style_lower
    ):

        return "List"

    # --------------------------------------------------------
    # REFERENCES
    # --------------------------------------------------------

    if (
        "reference" in style_lower
        or "bibliography" in style_lower
    ):

        return "Reference"

    return None


# ============================================================
# TEXT + FORMAT CLASSIFICATION
# ============================================================

def classify_by_text(
    text: str,
    paragraph
) -> Optional[str]:
    """
    Secondary classification when Word style
    is not sufficient.
    """

    if not text:

        return None

    words = text.split()

    # --------------------------------------------------------
    # FIGURE CAPTION
    # --------------------------------------------------------

    if starts_with_figure(text):

        return "Figure Caption"

    # --------------------------------------------------------
    # TABLE CAPTION
    # --------------------------------------------------------

    if starts_with_table(text):

        return "Table Caption"

    # --------------------------------------------------------
    # REFERENCE
    # --------------------------------------------------------

    if looks_like_reference(text):

        return "Reference"

    # --------------------------------------------------------
    # CHAPTER
    # --------------------------------------------------------

    if starts_with_chapter(text):

        return "Chapter Title"

    # --------------------------------------------------------
    # LIST
    # --------------------------------------------------------

    list_patterns = [

        # Bullet
        r"^\s*[-•▪◦]\s+",

        # Numbered list
        r"^\s*\d+[\.\)]\s+",

        # Alphabetical list
        r"^\s*[A-Za-z][\.\)]\s+",

    ]

    for pattern in list_patterns:

        if re.match(
            pattern,
            text
        ):

            return "List"

    # --------------------------------------------------------
    # IMAGE
    # --------------------------------------------------------

    if paragraph_contains_image(
        paragraph
    ):

        return "Figure"

    # --------------------------------------------------------
    # ALL CAPS SHORT TEXT
    # --------------------------------------------------------

    if (
        len(words) <= 12
        and calculate_uppercase_ratio(text)
        >= 0.75
    ):

        return "Heading 1"

    # --------------------------------------------------------
    # SHORT BOLD TEXT
    # --------------------------------------------------------

    if (
        get_bold(paragraph)
        and len(words) <= 12
    ):

        font_size = get_font_size(
            paragraph
        )

        if font_size >= 16:

            return "Heading 1"

        if font_size >= 13:

            return "Heading 2"

    # --------------------------------------------------------
    # DEFAULT
    # --------------------------------------------------------

    return "Paragraph"


# ============================================================
# TABLE EXTRACTION
# ============================================================

def extract_tables(
    document,
    source_document: str
):
    """
    Extract each Word table as one training sample.
    """

    samples = []

    for table in document.tables:

        rows = []

        for row in table.rows:

            cells = []

            for cell in row.cells:

                cell_text = clean_text(
                    cell.text
                )

                if cell_text:

                    cells.append(
                        cell_text
                    )

            if cells:

                rows.append(
                    " | ".join(cells)
                )

        table_text = " || ".join(
            rows
        )

        if not table_text:

            table_text = "[TABLE]"

        sample = {

            "text":
                table_text,

            "label":
                "Table",

            "source_document":
                source_document,

            "style":
                "Table",

            "style_id":
                "",

            "font_size":
                0,

            "bold":
                0,

            "italic":
                0,

            "underline":
                0,

            "alignment":
                "unknown",

            "left_indent":
                0,

            "right_indent":
                0,

            "first_line_indent":
                0,

            "space_before":
                0,

            "space_after":
                0,

            "line_spacing":
                0,

            "paragraph_index":
                -1,

            "text_length":
                len(table_text),

            "word_count":
                len(table_text.split()),

            "uppercase_ratio":
                calculate_uppercase_ratio(
                    table_text
                ),

            "digit_ratio":
                calculate_digit_ratio(
                    table_text
                ),

            "starts_with_number":
                0,

            "starts_with_chapter":
                0,

            "starts_with_figure":
                0,

            "starts_with_table":
                1,

            "looks_like_reference":
                0,

            "has_image":
                0,

            "has_table":
                1,
        }

        samples.append(
            sample
        )

    return samples


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_document(
    document_path: Path
):
    """
    Extract training samples from one DOCX.
    """

    print()
    print("----------------------------------------")
    print(
        f"Processing: {document_path.name}"
    )
    print("----------------------------------------")

    try:

        document = Document(
            str(document_path)
        )

    except Exception as error:

        print(
            f"ERROR opening document: {error}"
        )

        return []

    samples = []

    # --------------------------------------------------------
    # PARAGRAPHS
    # --------------------------------------------------------

    for index, paragraph in enumerate(
        document.paragraphs
    ):

        text = clean_text(
            paragraph.text
        )

        has_image = (
            paragraph_contains_image(
                paragraph
            )
        )

        # Ignore completely empty paragraphs
        # unless they contain an image.
        if not text and not has_image:

            continue

        if not text and has_image:

            text = "[IMAGE]"

        # ----------------------------------------------------
        # WORD STYLE
        # ----------------------------------------------------

        try:

            style_name = (
                paragraph.style.name
                or ""
            ).strip()

        except Exception:

            style_name = ""

        try:

            style_id = (
                paragraph.style.style_id
                or ""
            )

        except Exception:

            style_id = ""

        # ----------------------------------------------------
        # LABEL
        # ----------------------------------------------------

        label = classify_by_word_style(
            paragraph
        )

        if label is None:

            label = classify_by_text(
                text,
                paragraph
            )

        # Safety
        if label not in VALID_LABELS:

            label = "Paragraph"

        # ----------------------------------------------------
        # FEATURES
        # ----------------------------------------------------

        sample = {

            "text":
                text,

            "label":
                label,

            "source_document":
                document_path.name,

            "style":
                style_name,

            "style_id":
                style_id,

            "font_size":
                get_font_size(
                    paragraph
                ),

            "bold":
                get_bold(
                    paragraph
                ),

            "italic":
                get_italic(
                    paragraph
                ),

            "underline":
                get_underline(
                    paragraph
                ),

            "alignment":
                get_alignment(
                    paragraph
                ),

            "left_indent":
                get_indent(
                    paragraph,
                    "left_indent"
                ),

            "right_indent":
                get_indent(
                    paragraph,
                    "right_indent"
                ),

            "first_line_indent":
                get_indent(
                    paragraph,
                    "first_line_indent"
                ),

            "space_before":
                get_spacing(
                    paragraph,
                    "space_before"
                ),

            "space_after":
                get_spacing(
                    paragraph,
                    "space_after"
                ),

            "line_spacing":
                get_line_spacing(
                    paragraph
                ),

            "paragraph_index":
                index,

            "text_length":
                len(text),

            "word_count":
                len(text.split()),

            "uppercase_ratio":
                calculate_uppercase_ratio(
                    text
                ),

            "digit_ratio":
                calculate_digit_ratio(
                    text
                ),

            "starts_with_number":
                starts_with_number(
                    text
                ),

            "starts_with_chapter":
                starts_with_chapter(
                    text
                ),

            "starts_with_figure":
                starts_with_figure(
                    text
                ),

            "starts_with_table":
                starts_with_table(
                    text
                ),

            "looks_like_reference":
                looks_like_reference(
                    text
                ),

            "has_image":
                int(has_image),

            "has_table":
                0,
        }

        samples.append(
            sample
        )

    # --------------------------------------------------------
    # TABLES
    # --------------------------------------------------------

    table_samples = extract_tables(
        document,
        document_path.name
    )

    samples.extend(
        table_samples
    )

    print(
        f"Extracted samples: {len(samples)}"
    )

    return samples


# ============================================================
# REMOVE DUPLICATES
# ============================================================

def remove_duplicates(
    samples
):
    """
    Remove exact duplicate text + label combinations.
    """

    unique = {}

    for sample in samples:

        key = (
            sample["text"]
            .strip()
            .lower(),

            sample["label"],
        )

        if key not in unique:

            unique[key] = sample

    return list(
        unique.values()
    )


# ============================================================
# DATASET STATISTICS
# ============================================================

def print_statistics(
    samples
):

    print()
    print("========================================")
    print("DATASET STATISTICS")
    print("========================================")

    print(
        f"Total samples : {len(samples)}"
    )

    labels = {}

    for sample in samples:

        label = sample["label"]

        labels[label] = (
            labels.get(
                label,
                0
            )
            + 1
        )

    print()
    print("Samples per label:")
    print()

    for label in VALID_LABELS:

        count = labels.get(
            label,
            0
        )

        print(
            f"{label:<20} {count}"
        )

    print()
    print("----------------------------------------")

    # Documents contributing samples
    documents = set(
        sample["source_document"]
        for sample in samples
    )

    print(
        f"Source documents : {len(documents)}"
    )


# ============================================================
# DATASET QUALITY CHECK
# ============================================================

def quality_check(
    samples
):
    """
    Basic dataset quality check.

    This does NOT modify the dataset.
    """

    print()
    print("========================================")
    print("DATASET QUALITY CHECK")
    print("========================================")

    if not samples:

        print(
            "ERROR: Dataset contains no samples."
        )

        return False

    label_counts = {}

    for sample in samples:

        label = sample["label"]

        label_counts[label] = (
            label_counts.get(
                label,
                0
            )
            + 1
        )

    warnings = 0

    for label in VALID_LABELS:

        count = label_counts.get(
            label,
            0
        )

        if count == 0:

            print(
                f"WARNING: No samples for '{label}'"
            )

            warnings += 1

        elif count < 3:

            print(
                f"WARNING: Very few samples for "
                f"'{label}': {count}"
            )

            warnings += 1

    if warnings == 0:

        print(
            "Dataset quality check: PASSED"
        )

    else:

        print()
        print(
            f"Dataset quality warnings: {warnings}"
        )

        print(
            "More properly formatted documents "
            "are recommended."
        )

    return True


# ============================================================
# SAVE DATASET
# ============================================================

def save_dataset(
    samples
):

    DATASET_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    with open(
        OUTPUT_FILE,
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=CSV_FIELDS
        )

        writer.writeheader()

        for sample in samples:

            writer.writerow(
                sample
            )

    print()
    print("========================================")
    print("DATASET SAVED")
    print("========================================")

    print(
        f"Location:\n{OUTPUT_FILE}"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("========================================")
    print("DocuForge AI")
    print("Offline Dataset Generator")
    print("========================================")

    print()
    print(
        "Input directory:"
    )
    print(
        INPUT_DIR
    )

    print()
    print(
        "Output dataset:"
    )
    print(
        OUTPUT_FILE
    )

    # --------------------------------------------------------
    # CREATE DIRECTORIES
    # --------------------------------------------------------

    INPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    DATASET_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    # --------------------------------------------------------
    # FIND DOCX FILES
    # --------------------------------------------------------

    documents = sorted(
        INPUT_DIR.glob("*.docx")
    )

    print()
    print(
        f"Training documents found: "
        f"{len(documents)}"
    )

    if not documents:

        print()
        print("WARNING:")
        print(
            "No DOCX files found."
        )

        print()
        print(
            "Put properly formatted DOCX files here:"
        )

        print(
            INPUT_DIR
        )

        print()
        print("Example:")
        print(
            "research_paper.docx"
        )
        print(
            "thesis.docx"
        )
        print(
            "book.docx"
        )

        return

    # --------------------------------------------------------
    # EXTRACT ALL DOCUMENTS
    # --------------------------------------------------------

    all_samples = []

    for document_path in documents:

        samples = extract_document(
            document_path
        )

        all_samples.extend(
            samples
        )

    # --------------------------------------------------------
    # CHECK
    # --------------------------------------------------------

    if not all_samples:

        print()
        print(
            "ERROR:"
        )

        print(
            "No training samples were extracted."
        )

        return

    # --------------------------------------------------------
    # REMOVE DUPLICATES
    # --------------------------------------------------------

    print()
    print(
        "Removing duplicate samples..."
    )

    before = len(
        all_samples
    )

    all_samples = remove_duplicates(
        all_samples
    )

    after = len(
        all_samples
    )

    print(
        f"Before duplicates removal : {before}"
    )

    print(
        f"After duplicates removal  : {after}"
    )

    # --------------------------------------------------------
    # STATISTICS
    # --------------------------------------------------------

    print_statistics(
        all_samples
    )

    # --------------------------------------------------------
    # QUALITY CHECK
    # --------------------------------------------------------

    quality_check(
        all_samples
    )

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    save_dataset(
        all_samples
    )

    # --------------------------------------------------------
    # COMPLETE
    # --------------------------------------------------------

    print()
    print("========================================")
    print("GENERATION COMPLETE")
    print("========================================")

    print()
    print(
        "Next step:"
    )

    print(
        "Train the offline ML classifier using:"
    )

    print()

    print(
        r"python ml\train_model.py"
    )

    print()


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":

    main()