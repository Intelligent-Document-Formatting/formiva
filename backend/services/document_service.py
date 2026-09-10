from pathlib import Path
from typing import Dict, Any, List
from docx import Document


def extract_document_content(file_path: str) -> Dict[str, Any]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Document not found: {file_path}")

    document = Document(str(path))
    paragraphs: List[Dict[str, Any]] = []

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue

        paragraphs.append({
            "index": index,
            "text": text,
            "style": paragraph.style.name if paragraph.style else "Normal",
            "bold": any(run.bold is True for run in paragraph.runs),
            "italic": any(run.italic is True for run in paragraph.runs),
        })

    tables: List[Dict[str, Any]] = []
    for table_index, table in enumerate(document.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({
            "index": table_index,
            "rows": rows
        })

    return {
        "file_name": path.name,
        "paragraphs": paragraphs,
        "tables": tables,
        "sections": len(document.sections),
    }