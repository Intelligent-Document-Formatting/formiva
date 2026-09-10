from pathlib import Path
from typing import Dict, Any, Tuple
from docx import Document

ALLOWED_EXTENSIONS = {".docx"}


def validate_file_extension(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def validate_docx_file(file_path: str) -> Tuple[bool, str, Dict[str, Any]]:
    path = Path(file_path)

    if not path.exists():
        return False, "File does not exist on disk.", {}

    if not validate_file_extension(path.name):
        return False, "Unsupported file format. Only .docx manuscripts are accepted.", {}

    try:
        doc = Document(str(path))
        para_count = sum(1 for p in doc.paragraphs if p.text.strip())
        table_count = len(doc.tables)

        if para_count == 0 and table_count == 0:
            return False, "The DOCX document is empty.", {"paragraphs": 0, "tables": 0}

        metadata = {
            "paragraphs": para_count,
            "tables": table_count,
            "sections": len(doc.sections),
            "file_size_bytes": path.stat().st_size
        }
        return True, "Document verified successfully.", metadata
    except Exception as exc:
        return False, f"Document structure corrupted: {str(exc)}", {}


def validate_document(file_path: str) -> Dict[str, Any]:
    valid, message, metadata = validate_docx_file(file_path)
    return {
        "valid": valid,
        "message": message,
        "file_name": Path(file_path).name,
        "details": metadata
    }